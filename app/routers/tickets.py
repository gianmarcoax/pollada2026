import csv
import io
from datetime import datetime
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form, status
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Ticket, Evento, Usuario
from app.schemas import (
    VentaMultipleCreate,
    TicketOut,
    ConfirmarEntregaPaymentRequest,
    TicketUpdate
)
from app.auth import get_current_user

router = APIRouter(prefix="/tickets", tags=["Tickets y Ventas"])

# ─────────────── Helpers de parsing ───────────────

def _parse_float(val, default=0.0):
    """Convierte un valor a float limpiando símbolos de moneda."""
    try:
        return float(str(val).replace("S/", "").replace(",", "").strip())
    except Exception:
        return default

def _parse_bool(val):
    """Interpreta 'Sí', 'Si', 'True', '1' como True."""
    return str(val).strip().lower() in ("sí", "si", "true", "1", "yes")

def _parse_estado(val):
    """Normaliza el valor de estado al formato interno."""
    v = str(val).strip().lower()
    if "parcial" in v:
        return "parcialmente_pagado"
    if "pagado" in v or "pago" in v:
        return "pagado"
    return "separado"

def _parse_metodo(val):
    """Normaliza el método de pago."""
    v = str(val).strip().lower()
    if v in ("yape",):   return "yape"
    if v in ("plin",):   return "plin"
    if v in ("efectivo",): return "efectivo"
    return "ninguno"

def _parse_fecha(val):
    """Intenta parsear una fecha desde string."""
    if not val or str(val).strip() in ("", "-", "None"):
        return None
    try:
        return datetime.strptime(str(val).strip(), "%Y-%m-%d %H:%M:%S")
    except Exception:
        return None

def _rows_from_file(content: bytes, filename: str) -> list[dict]:
    """
    Lee un archivo Excel (.xlsx) o CSV y devuelve lista de dicts
    usando las columnas del formato de exportación estándar.
    Columnas esperadas (en ese orden):
      Nº Boleto | Código / DNI | Nombre Comprador | Carrera | Ciclo |
      Persona que Recoge | Estado | Monto Total (S/) | Monto Pagado (S/) |
      Monto Pendiente (S/) | Método de Pago | Entregado | Fecha de Entrega
    """
    rows = []
    ext = filename.lower().rsplit(".", 1)[-1]

    if ext in ("xlsx", "xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            rows.append(dict(zip(headers, [v for v in row])))

    elif ext == "csv":
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            rows.append(dict(row))
    else:
        raise ValueError(f"Formato de archivo no soportado: .{ext} (usa .xlsx o .csv)")

    return rows

def _row_to_ticket_data(row: dict, evento_id: int) -> dict:
    """Convierte una fila del archivo en un dict con los campos del modelo Ticket."""
    # Mapeo flexible de nombres de columna (soporta variaciones de nombre)
    def get(keys):
        for k in keys:
            for rk in row:
                if str(rk).strip().lower() == k.lower():
                    v = row[rk]
                    return v if v is not None else ""
        return ""

    numero = int(_parse_float(get(["nº boleto", "n° boleto", "numero_boleto", "boleto"]), 0))
    if numero <= 0:
        return None  # fila inválida

    monto_total  = _parse_float(get(["monto total (s/)", "monto_total", "total"]), 15.0)
    monto_pagado = _parse_float(get(["monto pagado (s/)", "monto_pagado", "pagado"]), 0.0)
    monto_pend   = max(0.0, round(monto_total - monto_pagado, 2))
    entregado    = _parse_bool(get(["entregado"]))
    fecha_entrega = _parse_fecha(get(["fecha de entrega", "fecha_hora_entrega"]))

    return dict(
        numero_boleto    = numero,
        evento_id        = evento_id,
        codigo_alumno    = str(get(["código / dni", "codigo / dni", "codigo_alumno", "código"])).strip() or "000000",
        nombre_alumno    = str(get(["nombre comprador", "nombre_alumno", "nombre"])).strip() or "Sin nombre",
        carrera          = str(get(["carrera"])).strip() or "INGENIERIA DE SISTEMAS",
        ciclo            = str(get(["ciclo"])).strip() or "1",
        nombre_recolector= str(get(["persona que recoge", "nombre_recolector", "recolector"])).strip() or None,
        estado           = _parse_estado(get(["estado"])),
        precio_unitario  = monto_total,
        monto_total      = monto_total,
        monto_pagado     = min(monto_pagado, monto_total),
        monto_pendiente  = monto_pend,
        metodo_pago      = _parse_metodo(get(["método de pago", "metodo de pago", "metodo_pago", "método"])),
        entregado        = entregado,
        fecha_hora_entrega = fecha_entrega,
    )


@router.post("/importar")
def importar_tickets(
    evento_id: int = Form(...),
    modo: str = Form(..., description="'merge' = combinar | 'replace' = reemplazar todo"),
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user)
):
    """
    Importa boletos desde un archivo Excel (.xlsx) o CSV al evento indicado.

    **Modos:**
    - `merge`:   Agrega boletos nuevos e actualiza los existentes (por número de boleto).
                 No elimina boletos que no estén en el archivo.
    - `replace`: Elimina TODOS los boletos del evento y los recrea desde el archivo.

    El archivo debe tener el mismo formato que el export del sistema.
    Requiere autenticación.
    """
    evento = db.query(Evento).filter(Evento.id == evento_id).first()
    if not evento:
        raise HTTPException(status_code=404, detail="Evento no encontrado")

    if modo not in ("merge", "replace"):
        raise HTTPException(status_code=400, detail="Modo inválido. Usa 'merge' o 'replace'.")

    # Leer el archivo
    try:
        content = archivo.file.read()
        rows = _rows_from_file(content, archivo.filename)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error al leer el archivo: {str(e)}")

    if not rows:
        raise HTTPException(status_code=400, detail="El archivo no contiene filas de datos.")

    # Parsear todas las filas
    ticket_data_list = []
    filas_invalidas = 0
    for row in rows:
        data = _row_to_ticket_data(row, evento_id)
        if data is None:
            filas_invalidas += 1
            continue
        ticket_data_list.append(data)

    if not ticket_data_list:
        raise HTTPException(status_code=400, detail="No se encontraron filas válidas con número de boleto en el archivo.")

    creados = 0
    actualizados = 0

    try:
        if modo == "replace":
            # Borrar todos los boletos del evento
            db.query(Ticket).filter(Ticket.evento_id == evento_id).delete()
            db.flush()

            # Insertar todos desde el archivo
            for data in ticket_data_list:
                ticket = Ticket(**data)
                db.add(ticket)
                creados += 1

        else:  # merge
            # Obtener boletos existentes indexados por numero_boleto
            existentes = {
                t.numero_boleto: t
                for t in db.query(Ticket).filter(Ticket.evento_id == evento_id).all()
            }

            for data in ticket_data_list:
                num = data["numero_boleto"]
                if num in existentes:
                    # Actualizar el existente
                    t = existentes[num]
                    for k, v in data.items():
                        if k not in ("id", "evento_id"):
                            setattr(t, k, v)
                    actualizados += 1
                else:
                    # Crear nuevo
                    ticket = Ticket(**data)
                    db.add(ticket)
                    creados += 1

        db.commit()

    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Error al guardar en la base de datos: {str(e)}")

    return {
        "ok": True,
        "modo": modo,
        "evento_id": evento_id,
        "total_filas_archivo": len(rows),
        "filas_invalidas": filas_invalidas,
        "creados": creados,
        "actualizados": actualizados,
        "mensaje": f"Importación completada: {creados} creados, {actualizados} actualizados."
    }


@router.post("/registrar-venta", response_model=List[TicketOut], status_code=status.HTTP_201_CREATED)
def registrar_venta_multiple(
    venta_in: VentaMultipleCreate,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user)
):
    """
    Registra la venta o separación de uno o varios boletos físicos numerados (hasta 20)
    asociados a un solo estudiante o código/DNI.
    Permite asignar la persona referencial que recogerá cada pollada individualmente.
    Requiere autenticación.
    """
    if not venta_in.boletos or len(venta_in.boletos) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Debe incluir al menos un número de boleto físico para registrar la venta."
        )
    
    if len(venta_in.boletos) > 20:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se puede vender más de 20 boletos por persona en una sola transacción."
        )

    evento = db.query(Evento).filter(Evento.id == venta_in.evento_id).first()
    if not evento:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="El evento especificado no existe"
        )
    
    if not evento.activo:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El evento especificado no está activo"
        )

    # Validar que los números de boletos físicos no estén ocupados en este evento
    numeros_solicitados = [b.numero_boleto for b in venta_in.boletos]
    if len(numeros_solicitados) != len(set(numeros_solicitados)):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No puede ingresar números de boletos físicos duplicados en la misma venta."
        )

    existentes = db.query(Ticket).filter(
        Ticket.evento_id == venta_in.evento_id,
        Ticket.numero_boleto.in_(numeros_solicitados)
    ).all()

    if existentes:
        ocupados_str = ", ".join(str(t.numero_boleto) for t in existentes)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Los siguientes números de boleto físico ya han sido vendidos o asignados anteriormente: #{ocupados_str}"
        )

    # Cálculos financieros por boleto
    total_boletos = len(venta_in.boletos)
    precio_unitario = venta_in.precio_unitario if venta_in.precio_unitario > 0 else 15.0
    
    monto_total_por_boleto = precio_unitario
    monto_pagado_por_boleto = venta_in.monto_pagado_total / total_boletos if total_boletos > 0 else 0.0
    
    # Ajustar por seguridad
    if monto_pagado_por_boleto > monto_total_por_boleto:
        monto_pagado_por_boleto = monto_total_por_boleto

    monto_pendiente_por_boleto = max(0.0, monto_total_por_boleto - monto_pagado_por_boleto)

    tickets_creados = []
    for item in venta_in.boletos:
        recolector = item.nombre_recolector.strip() if (item.nombre_recolector and item.nombre_recolector.strip()) else venta_in.nombre_alumno.strip()
        
        ticket = Ticket(
            numero_boleto=item.numero_boleto,
            evento_id=venta_in.evento_id,
            codigo_alumno=venta_in.codigo_alumno.strip(),
            nombre_alumno=venta_in.nombre_alumno.strip(),
            carrera=venta_in.carrera.strip() if venta_in.carrera else "INGENIERIA DE SISTEMAS",
            ciclo=venta_in.ciclo.strip() if venta_in.ciclo else "1",
            nombre_recolector=recolector,
            estado=venta_in.estado,
            precio_unitario=precio_unitario,
            monto_total=monto_total_por_boleto,
            monto_pagado=monto_pagado_por_boleto,
            monto_pendiente=monto_pendiente_por_boleto,
            metodo_pago=venta_in.metodo_pago or "ninguno",
            entregado=False,
            fecha_hora_entrega=None
        )
        db.add(ticket)
        tickets_creados.append(ticket)

    db.commit()
    for t in tickets_creados:
        db.refresh(t)

    return tickets_creados

@router.get("/buscar", response_model=List[TicketOut])
def buscar_tickets(
    numero_boleto: Optional[int] = Query(None, description="Filtrar por número de boleto físico"),
    codigo_alumno: Optional[str] = Query(None, description="Filtrar por código de alumno o DNI"),
    q: Optional[str] = Query(None, description="Búsqueda por número de boleto, código, comprador o persona que recoge"),
    evento_id: Optional[int] = Query(None, description="Filtrar por evento determinado"),
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user)
):
    """
    Busca tickets por número de boleto físico, código de alumno/DNI, nombre del comprador o persona que recoge.
    Devuelve todos los datos financieros (monto abonado, saldo pendiente, estado y entrega).
    Requiere autenticación.
    """
    query = db.query(Ticket)
    
    if evento_id is not None:
        query = query.filter(Ticket.evento_id == evento_id)

    if numero_boleto is not None:
        query = query.filter(Ticket.numero_boleto == numero_boleto)
    elif codigo_alumno:
        query = query.filter(Ticket.codigo_alumno == codigo_alumno.strip())
    elif q:
        term = q.strip()
        # Intentar convertir término a número por si buscaron "102"
        if term.isdigit():
            num_val = int(term)
            query = query.filter(
                or_(
                    Ticket.numero_boleto == num_val,
                    Ticket.codigo_alumno.ilike(f"%{term}%"),
                    Ticket.nombre_alumno.ilike(f"%{term}%"),
                    Ticket.nombre_recolector.ilike(f"%{term}%")
                )
            )
        else:
            query = query.filter(
                or_(
                    Ticket.codigo_alumno.ilike(f"%{term}%"),
                    Ticket.nombre_alumno.ilike(f"%{term}%"),
                    Ticket.nombre_recolector.ilike(f"%{term}%")
                )
            )
    
    return query.order_by(Ticket.numero_boleto.asc()).all()

@router.post("/confirmar-entrega", response_model=TicketOut)
def confirmar_entrega_boleto(
    datos: ConfirmarEntregaPaymentRequest,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user)
):
    """
    Confirma la entrega física de un boleto mediante su número impreso.
    Si el boleto ya fue entregado previamente, retorna HTTP 400 con la fecha exacta.
    Permite ingresar un cobro adicional en puerta si el ticket estaba parcialmente pagado o separado.
    Requiere autenticación.
    """
    ticket = db.query(Ticket).filter(Ticket.numero_boleto == datos.numero_boleto).first()
    if not ticket:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"No se encontró ningún boleto físico con el número #{datos.numero_boleto}."
        )

    if ticket.entregado:
        fecha_str = ticket.fecha_hora_entrega.strftime("%d/%m/%Y a las %H:%M:%S") if ticket.fecha_hora_entrega else "fecha desconocida"
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"El boleto físico #{datos.numero_boleto} ya fue entregado previamente el {fecha_str}."
        )

    # Si se pagó un saldo adicional durante la entrega
    if datos.monto_cobrado_adicional and datos.monto_cobrado_adicional > 0:
        ticket.monto_pagado += datos.monto_cobrado_adicional
        if ticket.monto_pagado >= ticket.monto_total:
            ticket.monto_pagado = ticket.monto_total
            ticket.monto_pendiente = 0.0
            ticket.estado = "pagado"
        else:
            ticket.monto_pendiente = max(0.0, ticket.monto_total - ticket.monto_pagado)
            ticket.estado = "parcialmente_pagado"

    if datos.metodo_pago_entrega and datos.metodo_pago_entrega != "ninguno":
        ticket.metodo_pago = datos.metodo_pago_entrega

    ticket.entregado = True
    ticket.fecha_hora_entrega = datetime.utcnow()

    db.commit()
    db.refresh(ticket)
    return ticket

@router.get("", response_model=List[TicketOut])
def listar_todos_tickets(
    evento_id: Optional[int] = None,
    estado: Optional[str] = None,
    entregado: Optional[bool] = None,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user)
):
    """
    Lista todos los tickets con filtros opcionales por evento, estado y entregado.
    """
    query = db.query(Ticket)
    if evento_id is not None:
        query = query.filter(Ticket.evento_id == evento_id)
    if estado is not None:
        query = query.filter(Ticket.estado == estado)
    if entregado is not None:
        query = query.filter(Ticket.entregado == entregado)

    return query.order_by(Ticket.numero_boleto.asc()).all()

@router.get("/{ticket_id}", response_model=TicketOut)
def obtener_ticket_por_id(
    ticket_id: int,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user)
):
    """
    Obtiene los datos de un ticket por su ID primario.
    """
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Ticket no encontrado"
        )
    return ticket

@router.patch("/{ticket_id}", response_model=TicketOut)
def editar_ticket(
    ticket_id: int,
    datos: TicketUpdate,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user)
):
    """
    Edita cualquier campo de un ticket existente: datos del comprador, recolector,
    estado de pago, montos, método de pago y estado de entrega.
    Recalcula automáticamente monto_pendiente al cambiar precio_unitario o monto_pagado.
    Requiere autenticación.
    """
    ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
    if not ticket:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"No se encontró ningún boleto con ID #{ticket_id}."
        )

    # Aplicar cambios de datos generales
    if datos.nombre_alumno is not None:
        ticket.nombre_alumno = datos.nombre_alumno.strip()
    if datos.codigo_alumno is not None:
        ticket.codigo_alumno = datos.codigo_alumno.strip()
    if datos.carrera is not None:
        ticket.carrera = datos.carrera.strip()
    if datos.ciclo is not None:
        ticket.ciclo = datos.ciclo.strip()
    if datos.nombre_recolector is not None:
        ticket.nombre_recolector = datos.nombre_recolector.strip()
    if datos.metodo_pago is not None:
        ticket.metodo_pago = datos.metodo_pago

    # Actualizar campos financieros y recalcular
    precio_nuevo = datos.precio_unitario if datos.precio_unitario is not None else ticket.precio_unitario
    pagado_nuevo = datos.monto_pagado if datos.monto_pagado is not None else ticket.monto_pagado

    if datos.precio_unitario is not None or datos.monto_pagado is not None:
        ticket.precio_unitario = precio_nuevo
        ticket.monto_total = precio_nuevo
        # Clamp: no puede pagar más de lo que vale
        ticket.monto_pagado = min(pagado_nuevo, precio_nuevo)
        ticket.monto_pendiente = max(0.0, precio_nuevo - ticket.monto_pagado)

        # Recalcular estado automáticamente si no se especificó uno nuevo
        if datos.estado is None:
            if ticket.monto_pendiente == 0.0:
                ticket.estado = "pagado"
            elif ticket.monto_pagado > 0.0:
                ticket.estado = "parcialmente_pagado"
            else:
                ticket.estado = "separado"

    # Estado explícito del usuario (tiene prioridad)
    if datos.estado is not None:
        ticket.estado = datos.estado

    # Estado de entrega
    if datos.entregado is not None:
        ticket.entregado = datos.entregado
        if datos.entregado and ticket.fecha_hora_entrega is None:
            ticket.fecha_hora_entrega = datetime.utcnow()
        elif not datos.entregado:
            ticket.fecha_hora_entrega = None

    # Fecha de entrega explícita (override manual)
    if datos.fecha_hora_entrega is not None:
        ticket.fecha_hora_entrega = datos.fecha_hora_entrega

    db.commit()
    db.refresh(ticket)
    return ticket
